"""
Conversión Markdown -> Word (.docx) nativa con python-docx.

Modernizado para ser determinista (recomendación de CONSTITUTION.md §7): la
conversión no depende de Pandoc, de modo que la salida .docx es idéntica en
cualquier entorno. Contexto histórico de los hallazgos de la Sección 5:

- 5.6: el fallback sin Pandoc convertía filas de tabla Markdown (`|...|`) en
  párrafos de texto plano, rompiendo el encabezado FT-SST-002 (una tabla).
  Hoy se construyen tablas Word reales con python-docx.
- 5.7: la rama Pandoc usaba como `reference_doc` el acta real con PII
  ("FT-SST-002 - Responsable del Sistema...docx"). Al eliminar la dependencia
  de Pandoc desapareció también el `reference_doc` y con él el riesgo de PII.
"""

import html
import os
import re

from docx.shared import Cm

from agente_sgsst.rendering.maquetador import _IMAGEN_MARKDOWN

_COLUMNA_SEPARADOR = re.compile(r"^:?-+:?$")

# Tablas HTML (las usa el informe del diagnóstico FT-SST-001, que maqueta la
# Tabla de Valores de la Res. 0312 como <table>: sin esto, se vertían a párrafos).
_RE_FILA_HTML = re.compile(r"<tr([^>]*)>(.*?)</tr>", re.DOTALL | re.IGNORECASE)
_RE_CELDA_HTML = re.compile(r"<t([dh])([^>]*)>(.*?)</t\1>", re.DOTALL | re.IGNORECASE)
_RE_COLSPAN = re.compile(r"colspan\s*=\s*['\"]?(\d+)", re.IGNORECASE)

ANCHO_IMAGEN_LOGO_CM = 2.8


def convertir_markdown_a_docx(md_path, docx_path):
    """
    Convierte un archivo Markdown a Word (.docx) usando el conversor nativo
    python-docx. No depende de Pandoc ni de ningún reference_doc: la salida es
    determinista en cualquier entorno.
    """
    if not os.path.exists(md_path):
        print(f"Error: Archivo Markdown no encontrado: {md_path}")
        return False

    return _convertir_con_python_docx(md_path, docx_path)


def _celdas_de_fila(linea: str) -> list[str]:
    """Divide una fila de tabla Markdown en sus celdas (sin pipes extremos)."""
    linea = linea.strip()
    if linea.startswith("|"):
        linea = linea[1:]
    if linea.endswith("|"):
        linea = linea[:-1]
    return [c.strip() for c in linea.split("|")]


def _es_fila_separadora(celdas: list[str]) -> bool:
    """Detecta la fila de separación `|---|---|` de Markdown."""
    return bool(celdas) and all(_COLUMNA_SEPARADOR.match(c) for c in celdas)


def _parsear_tabla(filas: list[str]) -> list[list[str]]:
    """Convierte líneas consecutivas `|...|` en una matriz de celdas."""
    tabla = []
    for fila in filas:
        celdas = _celdas_de_fila(fila)
        if _es_fila_separadora(celdas):
            continue
        tabla.append(celdas)
    return tabla


def _ruta_imagen_markdown(texto: str) -> str | None:
    """Si `texto` es una imagen Markdown `![alt](ruta)`, devuelve la ruta; si no, None."""
    match = _IMAGEN_MARKDOWN.match(texto.strip())
    if match:
        return match.group(2)
    return None


def _limpiar_celda_html(texto: str) -> str:
    """Extrae el texto plano de una celda HTML sin etiquetas ni entidades."""
    texto = re.sub(r"<br\s*/?>", " ", texto)
    texto = re.sub(r"<[^>]+>", "", texto)
    return html.unescape(texto).strip()


def _agregar_tabla_html(doc, lineas: list[str]) -> None:
    """Convierte un bloque `<table>...</table>` en una tabla Word real.

    Reutiliza la misma estrategia del hallazgo 5.6 pero para tablas HTML: cada
    `<tr>` es una fila, cada `<td>/<th>` una celda, respeta `colspan` (filas de
    encabezado/subtotales del formato Res. 0312) y resalta en amarillo las filas
    marcadas con `style="background-color: yellow;"` (hallazgo NA por justificar).
    """
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_COLOR_INDEX

    contenido = "\n".join(lineas)
    filas = []
    for m_fila in _RE_FILA_HTML.finditer(contenido):
        fila = {
            "amarillo": "yellow" in (m_fila.group(1) or "").lower(),
            "celdas": [],
        }
        for m_celda in _RE_CELDA_HTML.finditer(m_fila.group(2)):
            colspan = _RE_COLSPAN.search(m_celda.group(2))
            fila["celdas"].append(
                {
                    "texto": _limpiar_celda_html(m_celda.group(3)),
                    "span": int(colspan.group(1)) if colspan else 1,
                    "th": m_celda.group(1) == "h",
                }
            )
        if fila["celdas"]:
            filas.append(fila)

    if not filas:
        for linea in lineas:
            doc.add_paragraph(linea)
        return

    ncols = max(sum(c["span"] for c in fila["celdas"]) for fila in filas)
    tabla = doc.add_table(rows=len(filas), cols=ncols)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    try:
        tabla.style = "Table Grid"
    except Exception:
        pass

    for i_fila, fila in enumerate(filas):
        k_col = 0
        for celda in fila["celdas"]:
            span = celda["span"]
            celda_doc = tabla.rows[i_fila].cells[k_col]
            for i in range(1, span):
                if k_col + i < ncols:
                    celda_doc = celda_doc.merge(tabla.rows[i_fila].cells[k_col + i])
            parrafo = celda_doc.paragraphs[0]
            if span == 1 and celda["texto"]:
                ruta_imagen = _ruta_imagen_markdown(celda["texto"])
                if ruta_imagen is not None:
                    _insertar_imagen_en_celda(celda_doc, ruta_imagen)
                    k_col += span
                    continue
            run = parrafo.add_run(celda["texto"])
            if celda["th"] or i_fila == 0:
                run.bold = True
            if fila["amarillo"]:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            k_col += span


def _insertar_imagen_en_celda(celda, ruta_imagen: str):
    """Incrusta la imagen del logo en la celda de la cabecera FT-SST-002."""

    if not os.path.exists(ruta_imagen):
        celda.text = f"[Logo no encontrado: {ruta_imagen}]"
        return
    parrafo = celda.paragraphs[0]
    run = parrafo.add_run()
    try:
        run.add_picture(ruta_imagen, width=Cm(ANCHO_IMAGEN_LOGO_CM))
    except Exception as e:
        parrafo.text = f"[Logo no insertable: {str(e)}]"


def _convertir_con_python_docx(md_path, docx_path):
    """
    Conversor nativo con python-docx (único camino de conversión).

    Resultado del hallazgo 5.6: las líneas de tabla Markdown (`|...|`) ya NO se
    agregan como párrafos de texto plano; se construyen tablas Word reales
    (`doc.add_table`) con la primera fila en negrita.
    """
    try:
        import docx as docxlib
        from docx.enum.table import WD_TABLE_ALIGNMENT

        doc = docxlib.Document()

        with open(md_path, "r", encoding="utf-8") as f:
            lines = f.read().splitlines()

        n = 0
        while n < len(lines):
            linea = lines[n].strip()

            if not linea:
                n += 1
                continue

            if linea.startswith("# "):
                doc.add_heading(linea[2:], level=1)
            elif linea.startswith("## "):
                doc.add_heading(linea[3:], level=2)
            elif linea.startswith("### "):
                doc.add_heading(linea[4:], level=3)
            elif linea.lower().startswith("<table"):
                bloques = []
                while n < len(lines) and "</table>" not in lines[n]:
                    bloques.append(lines[n])
                    n += 1
                if n < len(lines):
                    bloques.append(lines[n])
                    n += 1
                _agregar_tabla_html(doc, bloques)
                continue
            elif linea.startswith("- ") or linea.startswith("* "):
                doc.add_paragraph(linea[2:].strip(), style="List Bullet")
            elif linea.startswith("|") or linea.endswith("|"):
                # Agrupar todas las líneas consecutivas de tabla
                filas = []
                while n < len(lines) and (lines[n].strip().startswith("|") or lines[n].strip().endswith("|")):
                    filas.append(lines[n])
                    n += 1
                tabla = _parsear_tabla(filas)
                if tabla:
                    ncols = max(len(r) for r in tabla)
                    tabla_doc = doc.add_table(rows=len(tabla), cols=ncols)
                    tabla_doc.alignment = WD_TABLE_ALIGNMENT.CENTER
                    try:
                        tabla_doc.style = "Table Grid"
                    except Exception:
                        pass
                    for i_fila, fila in enumerate(tabla):
                        for j_celda, celda in enumerate(fila):
                            parrafo = tabla_doc.rows[i_fila].cells[j_celda].paragraphs[0]
                            ruta_imagen = _ruta_imagen_markdown(celda)
                            if ruta_imagen is not None:
                                _insertar_imagen_en_celda(tabla_doc.rows[i_fila].cells[j_celda], ruta_imagen)
                                continue
                            run = parrafo.add_run(celda)
                            if i_fila == 0:
                                run.bold = True
                continue  # n ya avanzado dentro del agrupamiento
            else:
                ruta_imagen = _ruta_imagen_markdown(linea)
                if ruta_imagen is not None:
                    parrafo = doc.add_paragraph()
                    try:
                        parrafo.add_run().add_picture(ruta_imagen, width=Cm(ANCHO_IMAGEN_LOGO_CM))
                    except Exception as e:
                        parrafo.add_run(f"[Imagen no insertable: {str(e)}]")
                else:
                    doc.add_paragraph(linea)

            n += 1

        doc.save(docx_path)
        print(f"Convertido con éxito vía python-docx: {docx_path}")
        return True
    except Exception as e:
        print(f"Error crítico en conversión docx: {str(e)}")
        return False
