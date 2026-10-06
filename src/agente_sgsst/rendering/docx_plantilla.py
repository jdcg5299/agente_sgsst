"""
Generación de documentos Word (.docx) basada en la plantilla oficial FT-SST-002.
Copia `docs/plantilla_ft_sst_002.docx` (que conserva el encabezado de sección fusionado,
el pie de página de firmas, márgenes de 3 cm, tamaño Letter y fuente Times New Roman),
rellena dinámicamente las celdas del encabezado (código, estándar, nombre, fecha, empresa, logo),
actualiza el pie de página con el responsable y representante legal, y renderiza
el contenido Markdown generado por la IA en el cuerpo del documento.
"""

import os
from datetime import datetime
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.table import WD_TABLE_ALIGNMENT

_PLANTILLA_PATH = os.path.join(
    os.path.dirname(__file__),
    "..",
    "..",
    "..",
    "docs",
    "plantilla_ft_sst_002.docx",
)


def _set_texto_run(cell, nuevo_texto):
    """Reemplaza el texto de la primera celda manteniendo el formato del primer run."""
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = nuevo_texto
        # Eliminar runs extras para que el texto quede en un solo run
        for run in p.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        p.add_run(nuevo_texto)


def _llenar_encabezado_y_pie(section, doc_info, empresa, responsable_sst, fecha_str):
    """Rellena las tablas de encabezado y pie de página de la sección FT-SST-002
    basándose en la estructura fija de 6 filas x 4 columnas de la plantilla."""
    # 1. Encabezado (tabla de 6 filas x 4 cols en el header de la sección)
    if section.header.tables:
        table = section.header.tables[0]
        try:
            # R=0 C=1: celda del título institucional (se mantiene por defecto).
            # R=1 C=3: código de formato (p.ej. 'FT- SST - 002' -> FT-SST-012).
            _set_texto_run(table.rows[1].cells[3], doc_info["codigo"])
            # R=2 C=1: campo 'Estándar E1.1.1' -> 'Estándar E2.1.1'
            estandar = str(doc_info["estandar"]).strip()
            if not estandar.upper().startswith("E"):
                estandar = f"E{estandar}"
            _set_texto_run(table.rows[2].cells[1], f"Estándar {estandar}")
            # R=3 C=0: título del documento actual (celda fusionada vertical).
            _set_texto_run(table.rows[3].cells[0], doc_info["nombre"].upper())
            # R=3 C=3: fecha de generación del documento.
            _set_texto_run(table.rows[3].cells[3], fecha_str)
            # R=4 C=3 y R=5 C=3: versión y página (se conservan los valores de la plantilla).
        except Exception as e:
            print(f"Aviso: No se pudo rellenar encabezado ({e})")

        # Inserción de logo en la celda R=0 C=0 (celda vacía de la cabecera)
        logo_path = empresa.get("logo")
        if logo_path and os.path.exists(logo_path):
            try:
                cell_logo = table.rows[0].cells[0]
                cell_logo.text = ""  # limpiar posible texto placeholder
                p = cell_logo.paragraphs[0]
                run = p.add_run()
                run.add_picture(logo_path, width=Cm(2.5))
            except Exception as e:
                print(f"Aviso: No se pudo insertar logo en encabezado: {e}")

    # 2. Pie de página (Firmas: Elaboró, Revisó, Aprobó; Código)
    if section.footer.tables:
        table_f = section.footer.tables[0]
        nombre_resp = responsable_sst.get("nombre") or "Responsable SG-SST"
        cc_resp = responsable_sst.get("cc") or "C.C. ____________"
        rep_legal = empresa.get("representante_legal") or "Representante Legal"

        for row in table_f.rows:
            for cell in row.cells:
                p_text = cell.text
                if "Jean Frank" in p_text or ("Código:" in p_text and "FT-SST-" in p_text):
                    for p in cell.paragraphs:
                        if "Jean" in p.text:
                            _set_texto_run(cell, f"{nombre_resp}\n{cc_resp}")
                            break
                        elif "Código:" in p.text:
                            _set_texto_run(cell, f"Código: {doc_info['codigo']}")
                            break
                elif "Gladys" in p_text:
                    _set_texto_run(cell, f"{rep_legal}\nRepresentante Legal")


def _parsear_tabla_md(filas):
    """Parsea filas Markdown | col | col | en matriz limpia."""
    tabla = []
    for f in filas:
        f_str = f.strip()
        if f_str.startswith("|"):
            f_str = f_str[1:]
        if f_str.endswith("|"):
            f_str = f_str[:-1]
        celdas = [c.strip() for c in f_str.split("|")]
        # Ignorar línea separadora |---|---|
        if celdas and all(c.startswith("-") and c.endswith("-") for c in celdas if c):
            continue
        tabla.append(celdas)
    return [t for t in tabla if t and any(t)]


def _parsear_tabla_html(filas):
    """Parsea una tabla HTML plana (<tr><td>) a matriz limpia [][str]."""
    import re

    tabla = []
    for f in filas:
        f_str = f.strip()
        if "<tr>" not in f_str:
            continue
        celdas = [
            re.sub(r"<[^>]+>", "", c).strip() for c in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", f_str, re.S)
        ]
        if celdas:
            tabla.append(celdas)
    return tabla


def _agregar_tabla_word(doc, matriz):
    """Inserta una matriz como tabla Word con borde y Times New Roman."""
    if not matriz:
        return
    ncols = max(len(r) for r in matriz)
    t_doc = doc.add_table(rows=len(matriz), cols=ncols)
    t_doc.alignment = WD_TABLE_ALIGNMENT.CENTER
    try:
        t_doc.style = "Table Grid"
    except Exception:
        pass
    for i_r, fila in enumerate(matriz):
        for j_c, val in enumerate(fila):
            if j_c < len(t_doc.rows[i_r].cells):
                cell = t_doc.rows[i_r].cells[j_c]
                cell.text = val
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(11)
                        if i_r == 0:
                            r.bold = True


def generar_docx_ft_sst_002(md_path, docx_path, doc_info, contexto):
    """
    Genera un archivo Word (.docx) utilizando `docs/plantilla_ft_sst_002.docx`
    como base, rellenando el encabezado/pie institucional y parseando el
    contenido Markdown generado por la IA en el cuerpo del documento.
    """
    plantilla_real = os.path.abspath(_PLANTILLA_PATH)
    if not os.path.exists(plantilla_real):
        print(f"Error crítico: No se encontró la plantilla FT-SST-002 en {plantilla_real}")
        return False

    try:
        doc = Document(plantilla_real)
        empresa = contexto.get("empresa", {})
        responsable_sst = contexto.get("responsable_sst", {})
        fecha_str = datetime.now().strftime("%d/%m/%Y")

        # Rellenar encabezados y pies en todas las secciones
        for section in doc.sections:
            _llenar_encabezado_y_pie(section, doc_info, empresa, responsable_sst, fecha_str)

        # Leer contenido Markdown
        with open(md_path, "r", encoding="utf-8") as f:
            lines = f.read().splitlines()

        # Parsear e insertar en el cuerpo del documento
        n = 0
        while n < len(lines):
            linea = lines[n].strip()

            if not linea:
                n += 1
                continue

            # Ignorar cabeceras Markdown que ya están en el encabezado formal Word
            if "SISTEMA DE GESTIÓN" in linea.upper() or "FT-SST-" in linea or "ESTÁNDAR RES" in linea:
                n += 1
                continue

            if linea.startswith("# "):
                h = doc.add_paragraph()
                run = h.add_run(linea[2:].strip())
                run.bold = True
                run.font.size = Pt(16)
                run.font.name = "Times New Roman"
            elif linea.startswith("## "):
                h = doc.add_paragraph()
                run = h.add_run(linea[3:].strip())
                run.bold = True
                run.font.size = Pt(13)
                run.font.name = "Times New Roman"
            elif linea.startswith("### "):
                h = doc.add_paragraph()
                run = h.add_run(linea[4:].strip())
                run.bold = True
                run.font.size = Pt(12)
                run.font.name = "Times New Roman"
            elif linea.startswith("- ") or linea.startswith("* "):
                p = doc.add_paragraph()
                run = p.add_run(f"• {linea[2:].strip()}")
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)
            elif linea.startswith("<table>") or linea.startswith("<tr>") or linea.startswith("<table"):
                filas_html = []
                while n < len(lines) and not lines[n].strip().endswith("</table>"):
                    filas_html.append(lines[n])
                    n += 1
                if n < len(lines):
                    filas_html.append(lines[n])
                    n += 1
                matriz = _parsear_tabla_html(filas_html)
                _agregar_tabla_word(doc, matriz)
                continue
            elif linea.startswith("|") or linea.endswith("|"):
                # Agrupar filas de tabla
                filas = []
                while n < len(lines) and (lines[n].strip().startswith("|") or lines[n].strip().endswith("|")):
                    filas.append(lines[n])
                    n += 1
                matriz = _parsear_tabla_md(filas)
                # Omitir el bloque de encabezado institucional (ya está en el header formal Word)
                bloque = "\n".join(filas).upper()
                if "ESTÁNDAR RES" in bloque or "SISTEMA DE GESTIÓN" in bloque or "FT-SST-" in bloque:
                    continue
                if matriz:
                    _agregar_tabla_word(doc, matriz)
                continue
            else:
                p = doc.add_paragraph(linea)
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(12)

            n += 1

        doc.save(docx_path)
        print(f"Word generado con éxito (Plantilla FT-SST-002): {docx_path}")
        return True

    except Exception as e:
        print(f"Error generando docx con plantilla: {str(e)}")
        return False
