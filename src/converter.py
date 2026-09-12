"""
Conversión Markdown -> Word (.docx).

Corrige los hallazgos 5.6 y 5.7 de CONSTITUTION.md (Sección 5):
- 5.6: el fallback sin Pandoc convertía filas de tabla Markdown (`|...|`) en
  párrafos de texto plano, rompiendo el encabezado FT-SST-002 (una tabla).
  Ahora detecta y construye tablas Word reales con python-docx.
- 5.7: dejó de usarse por defecto el documento real con PII
  ("FT-SST-002 - Responsable del Sistema...docx") como reference_doc de Pandoc.
  El parámetro `reference_doc` es opcional y por defecto NO apunta a ningún
  documento con datos personales.
"""
import os
import re
import subprocess

_COLUMNA_SEPARADOR = re.compile(r"^:?-+:?$")

# El acta real con PII ya NO es el reference_doc por defecto (hallazgo 5.7).
# Un reference_doc solo aporta estilos y NO debe arrastrar datos de una persona.
REFERENCE_DOC_DEFAULT = None


def convertir_markdown_a_docx(md_path, docx_path, reference_doc=None):
    """
    Convierte un archivo Markdown a Word (.docx) utilizando Pandoc y la plantilla
    de referencia optional para heredar estilos corporativos. Si Pandoc no está
    instalado, utiliza un conversor nativo basado en python-docx (hallazgo 5.6).
    """
    if not os.path.exists(md_path):
        print(f"Error: Archivo Markdown no encontrado: {md_path}")
        return False

    if reference_doc is not None and "Responsable del Sistema de Gestión" in os.path.basename(reference_doc):
        print(
            "AVISO (hallazgo 5.7): se está usando como reference_doc el acta con PII de una "
            "persona real. No debería usarse como plantilla técnica; se continúa sin ella."
        )
        reference_doc = None

    # Intentar usar Pandoc si está disponible en el sistema
    try:
        cmd = ["pandoc", md_path, "-o", docx_path]
        if reference_doc and os.path.exists(reference_doc):
            cmd.extend(["--reference-doc", reference_doc])

        result = subprocess.run(cmd, capture_output=True, text=True)
        if result.returncode == 0:
            print(f"Convertido con éxito vía Pandoc: {docx_path}")
            return True
        print(f"Pandoc advertencia/fallo, usando fallback python-docx: {result.stderr}")
    except FileNotFoundError:
        print("Pandoc no detectado en el sistema, utilizando conversor nativo python-docx...")

    # Fallback nativo con python-docx
    return _convertir_con_python_docx(md_path, docx_path)


def _celdas_de_fila(linea: str) -> list[str]:
    """Divide una fila de tabla Markdown en sus celdas (sin pipes extremos)."""
    linea = linea.strip()
    if linea.startswith("|"):
        linea = linea[1:]
    if linea.endswith("|"):
        linea = linea[:-1]
    return [c.strip() for c in linea.split("|")]


def _es_fila_separadora(celdas: list[str]) -> bool:
    """Detecta la fila de separación `|---|---|` de Markdown."""
    return bool(celdas) and all(_COLUMNA_SEPARADOR.match(c) for c in celdas)


def _parsear_tabla(filas: list[str]) -> list[list[str]]:
    """Convierte líneas consecutivas `|...|` en una matriz de celdas."""
    tabla = []
    for fila in filas:
        celdas = _celdas_de_fila(fila)
        if _es_fila_separadora(celdas):
            continue
        tabla.append(celdas)
    return tabla


def _convertir_con_python_docx(md_path, docx_path):
    """
    Conversor nativo sin Pandoc usando python-docx.

    Corrige el hallazgo 5.6: las líneas de tabla Markdown (`|...|`) ya NO se
    agregan como párrafos de texto plano; se construyen tablas Word reales
    (`doc.add_table`) con la primera fila en negrita.
    """
    try:
        import docx as docxlib
        from docx.enum.table import WD_TABLE_ALIGNMENT

        doc = docxlib.Document()

        with open(md_path, "r", encoding="utf-8") as f:
            lines = f.read().splitlines()

        n = 0
        while n < len(lines):
            linea = lines[n].strip()

            if not linea:
                n += 1
                continue

            if linea.startswith("# "):
                doc.add_heading(linea[2:], level=1)
            elif linea.startswith("## "):
                doc.add_heading(linea[3:], level=2)
            elif linea.startswith("### "):
                doc.add_heading(linea[4:], level=3)
            elif linea.startswith("- ") or linea.startswith("* "):
                doc.add_paragraph(linea[2:].strip(), style="List Bullet")
            elif linea.startswith("|") or linea.endswith("|"):
                # Agrupar todas las líneas consecutivas de tabla
                filas = []
                while n < len(lines) and (lines[n].strip().startswith("|") or lines[n].strip().endswith("|")):
                    filas.append(lines[n])
                    n += 1
                tabla = _parsear_tabla(filas)
                if tabla:
                    ncols = max(len(r) for r in tabla)
                    tabla_doc = doc.add_table(rows=len(tabla), cols=ncols)
                    tabla_doc.alignment = WD_TABLE_ALIGNMENT.CENTER
                    try:
                        tabla_doc.style = "Table Grid"
                    except Exception:
                        pass
                    for i_fila, fila in enumerate(tabla):
                        for j_celda, celda in enumerate(fila):
                            parrafo = tabla_doc.rows[i_fila].cells[j_celda].paragraphs[0]
                            run = parrafo.add_run(celda)
                            if i_fila == 0:
                                run.bold = True
                continue  # n ya avanzado dentro del agrupamiento
            elif linea:
                doc.add_paragraph(linea)

            n += 1

        doc.save(docx_path)
        print(f"Convertido con éxito vía python-docx: {docx_path}")
        return True
    except Exception as e:
        print(f"Error crítico en conversión docx: {str(e)}")
        return False