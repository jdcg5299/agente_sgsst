"""
Tests del generador Excel del diagnóstico FT-SST-001 (entrega oficial .xlsx).

El diagnóstico se entrega en Excel replicando el instrumento oficial
(`docs/Diagnostico Resolucion 0312 de 2019 - 2026.xls`, hoja "Diagnostico
inicial"): 11 columnas — Ciclo | Numeral | Ítem | Criterio | Modo de verificación
| Valor del ítem del estándar | Peso porcentual | Puntaje Posible (Cumple
totalmente / No cumple / No aplica) | Calificación de la empresa — con la
jerarquía estándar → grupo → ítems del instrumento. Estos tests fijan ese
contrato para que una regresión no vuelva a entregar otro formato.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from agente_sgsst.domain.clasificacion import get_applicable_items
from agente_sgsst.domain.ponderacion import (
    bloques_evaluacion,
    calcular_diagnostico,
    cargar_estandares_capitulo_iii,
)
from agente_sgsst.rendering.excel_diagnostico import generar_excel_diagnostico

FILA_ENCABEZADO = 7
NOMBRE_HOJA = "Diagnostico inicial"
CAPITULO = "Capítulo II"
NCOLUMNAS = 11

# Índices 0-based de las columnas del instrumento oficial.
COL_CICLO, COL_NUMERAL, COL_ITEM, COL_CRITERIO, COL_MODO = 0, 1, 2, 3, 4
COL_VALOR, COL_PESO, COL_CUMPLE, COL_NOCUMPLE, COL_NOAPLICA, COL_CALIFICA = 5, 6, 7, 8, 9, 10

EMPRESA = {
    "razon_social": "ALEXA S.A.S.",
    "nit": "900123456",
    "total_trabajadores": 20,
    "clase_riesgo_arl": 1,
}


def _respuestas_aplicables(capitulo=CAPITULO, con=None):
    """Respuestas 'C' para todos los ítems APLICABLES del capítulo, con excepciones."""
    respuestas = {numeral: ("C", "") for numeral in get_applicable_items(capitulo)}
    if con:
        respuestas.update(con)
    return respuestas


def _generar(tmp_path: Path, respuestas, capitulo=CAPITULO, nombre="diagnostico.xlsx", empresa=None):
    ruta = tmp_path / nombre
    resultado = calcular_diagnostico(capitulo, respuestas) if respuestas else None
    generar_excel_diagnostico(capitulo, resultado, empresa or EMPRESA, ruta)
    return load_workbook(ruta)[NOMBRE_HOJA]


def _items() -> dict:
    return cargar_estandares_capitulo_iii()


def _filas_items(ws):
    """Mapea numeral -> fila para los ítems de la tabla (columna Numeral)."""
    items = _items()
    return {
        fila[COL_NUMERAL].value: fila
        for fila in ws.iter_rows(min_row=FILA_ENCABEZADO + 1)
        if fila[COL_NUMERAL].value in items
    }


def _valores(ws, columna):
    """Valores de una columna de todas las filas, como lista de strings ("" si vacío)."""
    return [str(ws.cell(row=r, column=columna + 1).value or "").strip() for r in range(1, ws.max_row + 1)]


def _es_amarillo(celda) -> bool:
    rgb = celda.fill.start_color.rgb
    return isinstance(rgb, str) and rgb.upper().endswith("FFFF00")


def _primer_aplicable() -> str:
    return sorted(get_applicable_items(CAPITULO))[0]


def _primero_no_aplicable() -> str:
    no_aplicables = set(_items()) - set(get_applicable_items(CAPITULO))
    return sorted(no_aplicables)[0]


class TestEstructuraOficial:
    def test_hoja_y_cabecera_de_la_empresa(self, tmp_path):
        ws = _generar(tmp_path, _respuestas_aplicables())
        assert ws.title == NOMBRE_HOJA
        assert "SISTEMA DE GESTIÓN DE LA SEGURIDAD Y SALUD EN EL TRABAJO" in ws["A1"].value
        assert "ALEXA S.A.S." in ws["A3"].value
        assert "900123456" in ws["A3"].value
        assert CAPITULO in ws["A4"].value

    def test_las_once_columnas_del_encabezado(self, tmp_path):
        ws = _generar(tmp_path, _respuestas_aplicables())
        titulos = [ws.cell(row=FILA_ENCABEZADO, column=c + 1).value for c in range(NCOLUMNAS)]
        assert titulos == [
            "Ciclo",
            "Numeral",
            "Item",
            "Criterio",
            "Modo de verificación",
            "Valor del item del estandar",
            "Peso porcentual",
            "Puntaje Posible",
            None,
            None,
            "Calificacion de la empresa o contratante",
        ]

    def test_subtitulos_del_puntaje_posible(self, tmp_path):
        ws = _generar(tmp_path, _respuestas_aplicables())
        fila = FILA_ENCABEZADO + 1
        assert [ws.cell(row=fila, column=c + 1).value for c in (COL_CUMPLE, COL_NOCUMPLE, COL_NOAPLICA)] == [
            "Cumple totalmente",
            "No cumple",
            "No aplica",
        ]

    def test_las_60_filas_de_items_estan_en_la_tabla(self, tmp_path):
        ws = _generar(tmp_path, _respuestas_aplicables())
        assert len(_filas_items(ws)) == 60


class TestTextoOficialPorItem:
    def test_cada_item_trae_item_criterio_y_modo_de_verificacion(self, tmp_path):
        items = _items()
        filas = _filas_items(_generar(tmp_path, _respuestas_aplicables()))
        for numeral, item in items.items():
            fila = filas[numeral]
            assert fila[COL_ITEM].value == item.descripcion
            assert fila[COL_CRITERIO].value == item.criterio
            assert fila[COL_MODO].value == item.modo_verificacion

    def test_el_texto_es_el_del_instrumento_y_no_un_resumen(self, tmp_path):
        filas = _filas_items(_generar(tmp_path, _respuestas_aplicables()))
        assert "diseñe e implemente" in filas["1.1.1"][COL_ITEM].value
        assert "Solicitar" in filas["1.1.1"][COL_MODO].value
        assert len(filas["4.2.3"][COL_CRITERIO].value) > 80


class TestJerarquiaEstandarGrupoItem:
    def test_encabezado_de_cada_estandar_con_su_porcentaje(self, tmp_path):
        ws = _generar(tmp_path, _respuestas_aplicables())
        primera = _valores(ws, COL_CICLO)
        for nombre, porcentaje in [
            ("RECURSOS", 10),
            ("GESTIÓN INTEGRAL DEL SISTEMA DE LA SEGURIDAD Y SALUD EN EL TRABAJO", 15),
            ("GESTIÓN DE LA SALUD", 20),
            ("GESTIÓN DE PELIGROS Y RIESGOS", 30),
            # El oficial escribe "GESTION" sin tilde: el encabezado se replica tal cual.
            ("GESTION DE AMENAZAS", 10),
            ("VERIFICACIÓN DEL SISTEMA DE GESTIÓN", 5),
            ("MEJORAMIENTO", 10),
        ]:
            assert any(nombre in v and f"({porcentaje}%)" in v for v in primera), nombre

    def test_encabezado_de_cada_grupo_que_el_instrumento_desglosa(self, tmp_path):
        """Solo los grupos con encabezado propio llevan fila propia: 1.1, 1.2, 3.x, 4.x.

        El instrumento NO desglosa los ítems de los estándares 2, 5, 6 y 7, así que
        inventarles un encabezado sería agregar una jerarquía que el oficial no tiene.
        """
        ws = _generar(tmp_path, _respuestas_aplicables())
        primera = _valores(ws, COL_CICLO)
        desglosados = [b for b in bloques_evaluacion() if b.desglosado]
        assert [b.clave for b in desglosados] == ["1.1", "1.2", "3.1", "3.2", "3.3", "4.1", "4.2"]
        for bloque in desglosados:
            esperado = f"{bloque.clave} {bloque.nombre} ({bloque.valor * 100:.0f}%)"
            assert esperado in primera, esperado

    def test_los_estandares_no_desglosados_no_inventan_encabezado_de_grupo(self, tmp_path):
        ws = _generar(tmp_path, _respuestas_aplicables())
        primera = _valores(ws, COL_CICLO)
        for bloque in [b for b in bloques_evaluacion() if not b.desglosado]:
            assert not any(v.startswith(f"{bloque.clave} {bloque.nombre}") for v in primera), bloque.clave

    def test_peso_porcentual_va_en_el_primer_item_de_cada_bloque(self, tmp_path):
        filas = _filas_items(_generar(tmp_path, _respuestas_aplicables()))
        for bloque in bloques_evaluacion():
            for indice, item in enumerate(bloque.items):
                esperado = bloque.valor if indice == 0 else None
                if esperado is None:
                    assert filas[item.numeral][COL_PESO].value is None, item.numeral
                else:
                    assert filas[item.numeral][COL_PESO].value == pytest.approx(esperado), item.numeral

    def test_el_porcentaje_total_cierra_cada_bloque_del_instrumento(self, tmp_path):
        """11 bloques: los 7 grupos desglosados más los 4 estándares sin desglose."""
        ws = _generar(tmp_path, _respuestas_aplicables())
        subtotales = [v for v in _valores(ws, COL_ITEM) if v == "PORCENTAJE TOTAL DEL ESTANDAR"]
        assert len(subtotales) == len(bloques_evaluacion()) == 11


class TestPuntajeCalculado:
    def test_todas_c_da_100_por_ciento(self, tmp_path):
        ws = _generar(tmp_path, _respuestas_aplicables())
        assert "100.00%" in ws["A4"].value
        assert "ACEPTABLE" in ws["A4"].value

    def test_no_cumple_iva_resta_del_puntaje(self, tmp_path):
        objetivo = _primer_aplicable()
        valor = _items()[objetivo].valor_item
        ws = _generar(tmp_path, _respuestas_aplicables(con={objetivo: ("NC", "")}))
        assert f"{(1.0 - valor) * 100:.2f}%" in ws["A4"].value

    def test_sin_respuestas_queda_pendiente_sin_inventar_puntaje(self, tmp_path):
        ws = _generar(tmp_path, None)
        assert "Pendiente de evaluación" in ws["A4"].value
        assert "—" in ws["A4"].value
        assert not any(fila[COL_CUMPLE].value for fila in _filas_items(ws).values())

    def test_columna_calificacion_suma_el_puntaje_del_motor(self, tmp_path):
        respuestas = _respuestas_aplicables()
        esperado = calcular_diagnostico(CAPITULO, respuestas).porcentaje
        ws = _generar(tmp_path, respuestas)
        suma = sum(fila[COL_CALIFICA].value for fila in _filas_items(ws).values())
        assert suma == pytest.approx(esperado)

    def test_suma_de_los_valores_de_los_items_es_100_por_ciento(self, tmp_path):
        ws = _generar(tmp_path, _respuestas_aplicables())
        suma = sum(fila[COL_VALOR].value for fila in _filas_items(ws).values())
        assert suma == pytest.approx(1.0)

    def test_cumple_se_registra_en_su_columna_y_suma_a_calificacion(self, tmp_path):
        numeral = _primer_aplicable()
        valor = _items()[numeral].valor_item
        fila = _filas_items(_generar(tmp_path, _respuestas_aplicables()))[numeral]
        assert fila[COL_CUMPLE].value == pytest.approx(valor)
        assert fila[COL_NOCUMPLE].value is None
        assert fila[COL_NOAPLICA].value is None
        assert fila[COL_CALIFICA].value == pytest.approx(valor)


class TestNoAplica:
    def test_no_aplica_automatico_no_se_resalta(self, tmp_path):
        fila = _filas_items(_generar(tmp_path, _respuestas_aplicables()))[_primero_no_aplicable()]
        assert fila[COL_NOAPLICA].value is not None
        assert not _es_amarillo(fila[COL_CICLO])

    def test_no_aplica_manual_se_resalta_en_amarillo(self, tmp_path):
        numeral = _primer_aplicable()
        respuestas = _respuestas_aplicables(con={numeral: ("NA", "No aplica para esta empresa.")})
        fila = _filas_items(_generar(tmp_path, respuestas))[numeral]
        assert _es_amarillo(fila[COL_CICLO])
        assert _es_amarillo(fila[COL_CALIFICA])

    def test_no_aplica_otorga_el_puntaje_completo_del_item(self, tmp_path):
        numeral = _primer_aplicable()
        respuestas = _respuestas_aplicables(con={numeral: ("NA", "No aplica para esta empresa.")})
        ws = _generar(tmp_path, respuestas)
        fila = _filas_items(ws)[numeral]
        assert fila[COL_NOAPLICA].value == pytest.approx(fila[COL_VALOR].value)
        assert "100.00%" in ws["A4"].value


class TestCierreDelLibro:
    def test_suma_total_por_estandar(self, tmp_path):
        ws = _generar(tmp_path, _respuestas_aplicables())
        assert _valores(ws, COL_ITEM).count("SUMA TOTAL") == 7

    def test_suma_total_de_los_estandares_minimos(self, tmp_path):
        ws = _generar(tmp_path, _respuestas_aplicables())
        fila = next(
            r
            for r in range(1, ws.max_row + 1)
            if str(ws.cell(row=r, column=COL_ITEM + 1).value or "").strip()
            == "SUMA TOTAL DE LOS ESTANDRES MINIMOS"
        )
        assert ws.cell(row=fila, column=COL_VALOR + 1).value == pytest.approx(1.0)

    def test_suma_total_de_los_avances_minimos_del_sg_sst(self, tmp_path):
        ws = _generar(tmp_path, _respuestas_aplicables())
        fila = next(
            r
            for r in range(1, ws.max_row + 1)
            if "SUMA TOTAL DE LOS AVANCES" in str(ws.cell(row=r, column=COL_ITEM + 1).value or "")
        )
        assert ws.cell(row=fila, column=COL_CALIFICA + 1).value == pytest.approx(1.0)

    def test_calificacion_final_con_nivel_del_articulo_27(self, tmp_path):
        ws = _generar(tmp_path, _respuestas_aplicables())
        ultima = str(ws.cell(row=ws.max_row, column=1).value or "")
        assert "CALIFICACIÓN DE LA EMPRESA" in ultima
        assert "100.00%" in ultima
        assert "ACEPTABLE" in ultima


class _RelojQueAvanza:
    """Reloj falso: cada llamada a `now()` devuelve la siguiente fecha de la lista."""

    def __init__(self, fechas):
        self._fechas = iter(fechas)

    def now(self):
        return datetime.strptime(next(self._fechas), "%d/%m/%Y")


class TestFechaDeGeneracion:
    """La fecha del encabezado se resuelve en cada llamada, no al importar el módulo.

    Si se calculara una sola vez al importar, dos diagnósticos generados en fechas
    distintas llevarían la misma, que es justo el defecto que se corrigió.
    """

    def test_dos_generaciones_en_dias_distintos_llevan_fecha_distinta(self, tmp_path, monkeypatch):
        import agente_sgsst.rendering.excel_diagnostico as modulo

        monkeypatch.setattr(modulo, "datetime", _RelojQueAvanza(["01/02/2026", "03/04/2026"]))
        primera = _generar(tmp_path, _respuestas_aplicables(), nombre="primero.xlsx")
        segunda = _generar(tmp_path, _respuestas_aplicables(), nombre="segundo.xlsx")
        assert "01/02/2026" in str(primera["H3"].value)
        assert "03/04/2026" in str(segunda["H3"].value)

    def test_el_encabezado_lleva_la_fecha_de_hoy(self, tmp_path):
        ws = _generar(tmp_path, _respuestas_aplicables())
        assert datetime.now().strftime("%d/%m/%Y") in str(ws["H3"].value)


class TestDatosDeEmpresaNoEjecutables:
    """La razón social y el NIT vienen de un formulario: no pueden volverse fórmulas.

    Excel interpreta como fórmula toda celda que empiece por =, +, - o @, así que
    "=HYPERLINK(...)" en el nombre de la empresa se ejecutaría al abrir el archivo.
    """

    def test_un_nombre_que_parece_formula_queda_texto(self, tmp_path):
        empresa = {**EMPRESA, "razon_social": '=HYPERLINK("http://ejemplo.test","clic")'}
        ws = _generar(tmp_path, _respuestas_aplicables(), empresa=empresa)
        celda = ws["A3"]
        assert not str(celda.value).lstrip("'").startswith("=")
        assert "HYPERLINK" in str(celda.value)
        assert celda.data_type != "f"

    @pytest.mark.parametrize("prefijo", ["=", "+", "-", "@"])
    def test_cualquier_prefijo_de_formula_queda_neutralizado(self, tmp_path, prefijo):
        empresa = {**EMPRESA, "razon_social": f"{prefijo}1+1", "nit": f"{prefijo}SUM(A1)"}
        ws = _generar(tmp_path, _respuestas_aplicables(), empresa=empresa)
        assert f"EMPRESA: '{prefijo}1+1" in str(ws["A3"].value)
        assert f"NIT '{prefijo}SUM(A1)" in str(ws["A3"].value)

    def test_un_nombre_normal_no_se_altera(self, tmp_path):
        ws = _generar(tmp_path, _respuestas_aplicables())
        assert "ALEXA S.A.S." in str(ws["A3"].value)
