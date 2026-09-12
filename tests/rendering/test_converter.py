"""
Tests del conversor Markdown -> docx (hallazgo 5.6 de CONSTITUTION.md).

Antes, el fallback sin Pandoc convertía las filas de tabla Markdown en párrafos
de texto plano (`doc.add_paragraph(line_str)`), rompiendo el encabezado FT-SST-002
en entornos sin Pandoc. Estos tests verifican que ahora se construye una tabla
Word real con python-docx.
"""
from __future__ import annotations

from pathlib import Path

from agente_sgsst.rendering.converter import _convertir_con_python_docx, _parsear_tabla


class TestParseoTabla:
    def test_ignora_fila_separadora(self):
        filas = [
            "| A | B |",
            "|---|---|",
            "| x | y |",
        ]
        matriz = _parsear_tabla(filas)
        assert matriz == [["A", "B"], ["x", "y"]]

    def test_recorta_pipes_extremos_y_espacios(self):
        matriz = _parsear_tabla(["| 1 |  2  |"])
        assert matriz == [["1", "2"]]


class TestConversionDocx:
    def test_tabla_markdown_se_convierte_en_tabla_word(self, tmp_path):
        md = tmp_path / "doc.md"
        md.write_text(
            "# Encabezado FT-SST-002\n\n"
            "| Logo | SISTEMA | CODIGO |\n"
            "| :---: | :--- | :---: |\n"
            "| | Titulo | FT-SST-001 |\n",
            encoding="utf-8",
        )
        out = tmp_path / "doc.docx"

        ok = _convertir_con_python_docx(str(md), str(out))
        assert ok is True
        assert out.exists()

        import docx

        document = docx.Document(str(out))
        assert len(document.tables) == 1
        tabla = document.tables[0]
        # 2 filas de datos (la separadora se descarta)
        assert len(tabla.rows) == 2
        assert len(tabla.columns) == 3
        assert tabla.rows[0].cells[0].text == "Logo"          # encabezado
        assert tabla.rows[1].cells[2].text == "FT-SST-001"     # dato

    def test_encabezado_markdown_se_mantiene_como_heading(self, tmp_path):
        md = tmp_path / "doc.md"
        md.write_text("# Diagnóstico Inicial\n\nPárrafo introductorio.\n", encoding="utf-8")
        out = tmp_path / "doc.docx"

        assert _convertir_con_python_docx(str(md), str(out)) is True

        import docx

        document = docx.Document(str(out))
        assert document.paragraphs[0].text == "Diagnóstico Inicial"