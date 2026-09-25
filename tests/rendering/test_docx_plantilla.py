import os
import tempfile
from docx import Document
from agente_sgsst.rendering.docx_plantilla import generar_docx_ft_sst_002, _parsear_tabla_html


def test_generar_docx_ft_sst_002():
    with tempfile.TemporaryDirectory() as tmpdir:
        md_path = os.path.join(tmpdir, "test.md")
        docx_path = os.path.join(tmpdir, "test.docx")

        with open(md_path, "w", encoding="utf-8") as f:
            f.write("# Título de Prueba\n\nEste es un párrafo de prueba.\n\n| Col 1 | Col 2 |\n|---|---|\n| Val 1 | Val 2 |\n")

        doc_info = {
            "codigo": "FT-SST-001",
            "estandar": "1.1.1",
            "nombre": "Prueba Doc",
        }
        contexto = {
            "empresa": {
                "razon_social": "Empresa Test SAS",
                "nit": "900123456-1",
                "representante_legal": "Juan Perez",
            },
            "responsable_sst": {
                "nombre": "Maria Gomez",
                "cc": "12345678",
            },
        }

        success = generar_docx_ft_sst_002(md_path, docx_path, doc_info, contexto)
        assert success is True
        assert os.path.exists(docx_path)
        assert os.path.getsize(docx_path) > 0


def test_parsear_tabla_html():
    filas = [
        "<table>",
        "  <tr><th>N°</th><th>Numeral</th></tr>",
        "  <tr><td>1</td><td>1.1.1</td></tr>",
        "</table>",
    ]
    matriz = _parsear_tabla_html(filas)
    assert matriz == [
        ["N\u00b0", "Numeral"],
        ["1", "1.1.1"],
    ]


def test_generar_docx_con_tabla_html():
    with tempfile.TemporaryDirectory() as tmpdir:
        md_path = os.path.join(tmpdir, "test_html.md")
        docx_path = os.path.join(tmpdir, "test_html.docx")

        with open(md_path, "w", encoding="utf-8") as f:
            f.write("<table>\n<tr><th>Numeral</th><th>Valor</th></tr>\n<tr><td>2.3.1</td><td>0.10</td></tr>\n</table>\n")

        doc_info = {"codigo": "FT-SST-001", "estandar": "E2.3.1", "nombre": "Diagnostico"}
        contexto = {
            "empresa": {"razon_social": "Empresa Test SAS", "nit": "900123456-1", "representante_legal": "Juan Perez"},
            "responsable_sst": {"nombre": "Maria Gomez", "cc": "12345678"},
        }

        success = generar_docx_ft_sst_002(md_path, docx_path, doc_info, contexto)
        assert success is True

        doc = Document(docx_path)
        assert len(doc.tables) == 1
        assert doc.tables[0].rows[0].cells[0].text == "Numeral"
        assert doc.tables[0].rows[1].cells[1].text == "0.10"
